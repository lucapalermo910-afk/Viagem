"""Gera public/modelo-roteiro.xlsx: o modelo vazio que o app importa e exporta.
Uso: python3 scripts/gerar-modelo.py
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

LINHAS = 301  # dados nas linhas 2 a 301
F = 'Arial'
OBRIG = PatternFill('solid', fgColor='E4572E')
OPC = PatternFill('solid', fgColor='6F665D')
AUTO = PatternFill('solid', fgColor='D9D2C7')
DATA, HORA, DIN = 'dd/mm/yyyy', 'hh:mm', '#,##0.00'

wb = Workbook()

def aba(nome, colunas, listas=None, formulas=None):
    """colunas: (título, largura, obrigatória, formato). formulas: {título: modelo com {r}}."""
    ws = wb.create_sheet(nome)
    for j, (titulo, larg, obrig, fmt) in enumerate(colunas, 1):
        c = ws.cell(1, j, titulo)
        auto = formulas and titulo in formulas
        c.font = Font(name=F, bold=True, color='2A2420' if auto else 'FFFFFF')
        c.fill = AUTO if auto else OBRIG if obrig else OPC
        c.alignment = Alignment(vertical='center', wrap_text=True)
        L = get_column_letter(j)
        ws.column_dimensions[L].width = larg
        for r in range(2, LINHAS + 1):
            cel = ws.cell(r, j)
            cel.font = Font(name=F, italic=bool(auto), color='8A5A00' if auto else '000000')
            if fmt: cel.number_format = fmt
            if auto: cel.value = formulas[titulo].format(r=r)
        if listas and titulo in listas:
            dv = DataValidation(type='list', formula1='"' + ','.join(listas[titulo]) + '"', allow_blank=True)
            dv.error, dv.errorTitle = 'Escolha um valor da lista.', 'Valor fora da lista'
            ws.add_data_validation(dv); dv.add(f'{L}2:{L}{LINHAS}')
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'A2'
    return ws

SIM_NAO = ['Sim', 'Não']
aba('Dias', [('Data', 13, True, DATA), ('Cidade', 18, False, None), ('Programa', 48, False, None), ('Notas', 48, False, None),
             ('Compromissos (automático)', 16, False, None), ('Pendência (automática)', 26, False, None)],
    formulas={'Compromissos (automático)': '=IF(A{r}="","",COUNTIF(Compromissos!A:A,A{r})+COUNTIF(Transporte!A:A,A{r}))',
              'Pendência (automática)': '=IF(A{r}="","",IF(AND(E{r}=0,C{r}=""),"Dia sem programa",""))'})

aba('Compromissos', [('Data', 13, True, DATA), ('Início', 9, False, HORA), ('Duração (min)', 11, False, '0'), ('Tipo', 14, True, None),
                     ('Nome', 30, True, None), ('Local', 34, False, None), ('Abre', 9, False, HORA), ('Fecha', 9, False, HORA),
                     ('Fechado no dia', 10, False, None), ('Trajeto (min)', 10, False, '0'), ('Reserva', 13, False, None),
                     ('Código', 18, False, None), ('Nome na reserva', 18, False, None), ('Pessoas', 9, False, '0'),
                     ('Pagamento', 16, False, None), ('Custo', 12, False, DIN), ('Pago', 8, False, None), ('Observações', 40, False, None),
                     ('Pendência (automática)', 44, False, None)],
    listas={'Tipo': ['Passeio', 'Restaurante', 'Ingresso', 'Compras', 'Livre'], 'Reserva': ['Sim', 'Não', 'Não precisa', 'A decidir'],
            'Fechado no dia': SIM_NAO, 'Pago': SIM_NAO},
    formulas={'Pendência (automática)':
        '=IF(E{r}="","",IF(AND(OR(D{r}="Restaurante",D{r}="Ingresso"),OR(K{r}="",K{r}="A decidir")),"Reserva a decidir; ","")'
        '&IF(AND(D{r}="Ingresso",K{r}="Não"),"Ingresso não comprado; ","")'
        '&IF(AND(K{r}="Sim",L{r}=""),"Reserva sem código; ","")'
        '&IF(I{r}="Sim","Fechado no dia; ","")'
        '&IF(AND(B{r}<>"",G{r}<>"",B{r}<G{r}),"Antes de abrir; ","")'
        '&IF(AND(B{r}<>"",H{r}<>"",B{r}+C{r}/1440>H{r}),"Depois de fechar; ",""))'})

aba('Hospedagem', [('Hotel', 34, True, None), ('Cidade', 16, False, None), ('Check-in', 13, True, DATA), ('Hora check-in', 10, False, HORA),
                   ('Check-out', 13, True, DATA), ('Hora check-out', 10, False, HORA), ('Endereço', 44, False, None), ('Link', 30, False, None),
                   ('Código', 18, False, None), ('Valor total', 13, False, DIN), ('Pago', 8, False, None), ('Observações', 40, False, None),
                   ('Pendência (automática)', 40, False, None)],
    listas={'Pago': SIM_NAO},
    formulas={'Pendência (automática)':
        '=IF(A{r}="","",IF(I{r}="","Sem código; ","")&IF(G{r}="","Sem endereço; ","")'
        '&IF(OR(C{r}="",E{r}=""),"Sem datas; ","")&IF(AND(C{r}<>"",E{r}<>"",E{r}<=C{r}),"Check-out antes do check-in; ",""))'})

aba('Transporte', [('Data', 13, True, DATA), ('Tipo', 11, True, None), ('De', 22, True, None), ('Para', 22, True, None),
                   ('Partida', 9, True, HORA), ('Chegada', 9, False, HORA), ('Número', 12, False, None), ('Assento', 10, False, None),
                   ('Código', 18, False, None), ('Terminal', 14, False, None), ('Custo', 12, False, DIN), ('Pago', 8, False, None),
                   ('Observações', 40, False, None), ('Pendência (automática)', 34, False, None)],
    listas={'Tipo': ['Voo', 'Trem', 'Ônibus', 'Outro'], 'Pago': SIM_NAO},
    formulas={'Pendência (automática)':
        '=IF(A{r}="","",IF(E{r}="","Sem horário de partida; ","")&IF(AND(B{r}="Voo",I{r}=""),"Voo sem código; ",""))'})

# ---- Pendências e custos (tudo por fórmula)
ws = wb.create_sheet('Pendências')
ws.column_dimensions['A'].width = 44; ws.column_dimensions['B'].width = 16; ws.column_dimensions['C'].width = 16; ws.column_dimensions['D'].width = 16
def titulo(r, t):
    c = ws.cell(r, 1, t); c.font = Font(name=F, bold=True, size=13, color='E4572E')
def linha(r, rotulo, *formulas, fmt=None):
    ws.cell(r, 1, rotulo).font = Font(name=F)
    for j, f in enumerate(formulas, 2):
        c = ws.cell(r, j, f); c.font = Font(name=F, bold=True)
        if fmt: c.number_format = fmt
titulo(1, 'Pendências (contagem automática)')
P = [('Restaurantes e ingressos com reserva a decidir', '=COUNTIF(Compromissos!S:S,"*Reserva a decidir*")'),
     ('Ingressos não comprados', '=COUNTIF(Compromissos!S:S,"*Ingresso não comprado*")'),
     ('Reservas sem código', '=COUNTIF(Compromissos!S:S,"*Reserva sem código*")'),
     ('Lugares fechados no dia da visita', '=COUNTIF(Compromissos!S:S,"*Fechado no dia*")'),
     ('Compromissos fora do horário de funcionamento', '=COUNTIF(Compromissos!S:S,"*Antes de abrir*")+COUNTIF(Compromissos!S:S,"*Depois de fechar*")'),
     ('Hospedagens com pendência', '=COUNTIF(Hospedagem!M2:M301,"?*")'),
     ('Transportes com pendência', '=COUNTIF(Transporte!N2:N301,"?*")'),
     ('Dias sem programa', '=COUNTIF(Dias!F:F,"Dia sem programa")')]
for i, (rot, f) in enumerate(P, 2): linha(i, rot, f)
linha(10, 'Total de pendências', '=SUM(B2:B9)')
titulo(12, 'Custos (na moeda em que você preencheu)')
for j, t in enumerate(['Total', 'Pago', 'A pagar'], 2):
    ws.cell(13, j, t).font = Font(name=F, bold=True, color='6F665D')
C = [('Compromissos', 'Compromissos!P:P', 'Compromissos!Q:Q'), ('Hospedagem', 'Hospedagem!J:J', 'Hospedagem!K:K'), ('Transporte', 'Transporte!K:K', 'Transporte!L:L')]
for i, (rot, val, pago) in enumerate(C, 14):
    linha(i, rot, f'=SUM({val})', f'=SUMIF({pago},"Sim",{val})', f'=B{i}-C{i}', fmt=DIN)
linha(17, 'Total', '=SUM(B14:B16)', '=SUM(C14:C16)', '=SUM(D14:D16)', fmt=DIN)

# ---- Leia-me
ws = wb['Sheet']; ws.title = 'Leia-me'
ws.column_dimensions['A'].width = 24; ws.column_dimensions['B'].width = 100
linhas = [
    ('Modelo de roteiro', 'Preencha as abas Dias, Compromissos, Hospedagem e Transporte. O app importa este arquivo e também exporta para ele.'),
    ('Cores do cabeçalho', 'Laranja: obrigatório. Cinza escuro: opcional. Cinza claro: automático (não preencha).'),
    ('Datas e horários', 'Datas como 24/10/2026. Horários como 17:30, no horário local do lugar.'),
    ('Dias', 'Opcional. Uma linha por dia, para informar a cidade e anotações. Exemplo: 24/10/2026 | Kanazawa | Jardim e bairro das gueixas | Levar guarda-chuva'),
    ('Compromissos', 'Uma linha por atividade. Exemplo: 24/10/2026 | 12:30 | 60 | Restaurante | Sushi Exemplo | Rua Fictícia 1, Kanazawa | 11:30 | 14:00 | | 15 | Sim | ABC123 | Silva | 2 | Só dinheiro | 180,00 | Não'),
    ('Tipo', 'Passeio, Restaurante, Ingresso (atração com entrada comprada), Compras ou Livre.'),
    ('Reserva', 'Sim, Não, Não precisa ou A decidir. Restaurante e Ingresso sem resposta contam como pendência.'),
    ('Abre / Fecha', 'Horário de funcionamento no dia da visita. Se o lugar não abre nesse dia, marque Sim em "Fechado no dia".'),
    ('Trajeto (min)', 'Tempo para chegar a este compromisso. O app usa para calcular a hora de sair.'),
    ('Hospedagem', 'Uma linha por hotel. Exemplo: Hotel Exemplo | Kanazawa | 23/10/2026 | 15:00 | 25/10/2026 | 11:00 | Rua Fictícia 2 | https://example.com | XYZ789 | 2.400,00 | Sim'),
    ('Transporte', 'Uma linha por trecho. Exemplo: 23/10/2026 | Trem | Tokyo | Kanazawa | 09:20 | 12:05 | Kagayaki 505 | 7A | | Plataforma 21 | 650,00 | Sim'),
    ('Pendências', 'Contagens e custos calculados sozinhos. Não edite.'),
    ('Limite', 'As fórmulas e listas cobrem até a linha 301 de cada aba.'),
]
for i, (a, b) in enumerate(linhas, 1):
    ws.cell(i, 1, a).font = Font(name=F, bold=True, size=14 if i == 1 else 11, color='E4572E' if i == 1 else '000000')
    c = ws.cell(i, 2, b); c.font = Font(name=F); c.alignment = Alignment(wrap_text=True, vertical='top')

wb.calculation.fullCalcOnLoad = True
wb.save('public/modelo-roteiro.xlsx')
print('ok', wb.sheetnames)
